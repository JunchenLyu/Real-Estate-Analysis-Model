# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt

PATH = r"C:\Users\whbjl\AppData\Local\Temp\codebuddy-dropped-files\b8337adf-5cb1-4da2-a1ce-fbb6fac1e070\参赛报告——房地产产业链风险传导分析金融智能体项目_v1.1_优化版.docx"
d = Document(PATH)
paras = d.paragraphs
print("总段落:", len(paras), "| 表格:", len(d.tables), "| 插图:", len(d.inline_shapes))
print()
for i, p in enumerate(paras[:40]):
    t = p.text.strip()
    if t:
        sz = p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else "-"
        print(i, "[", p.alignment, "]", "sz=", sz, "|", t[:50])